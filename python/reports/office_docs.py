from docx import Document
from pathlib import Path
def build_summary(title,lines,path='reports/ringkasan.docx'):
 Path(path).parent.mkdir(parents=True,exist_ok=True);doc=Document();doc.add_heading(title,0)
 for line in lines:doc.add_paragraph(str(line))
 doc.add_paragraph('PT NEXOVONARSACORPORATION - All Right Reserved');doc.save(path);return path
if __name__=='__main__':build_summary('Laporan Operasional',[])