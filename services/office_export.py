from pathlib import Path
from openpyxl import Workbook
from docx import Document
ROOT=Path("exports");ROOT.mkdir(exist_ok=True)
def xlsx(name,headers,rows):
    wb=Workbook();ws=wb.active;ws.title="Data";ws.append(headers)
    for row in rows: ws.append(row)
    out=ROOT/name;wb.save(out);return str(out)
def docx(name,title,paragraphs):
    doc=Document();doc.add_heading(title,0)
    for p in paragraphs:doc.add_paragraph(p)
    out=ROOT/name;doc.save(out);return str(out)
def generate_pack(data):
    return {
      "member_resmi":xlsx("member_resmi.xlsx",data["member_headers"],data["members"]),
      "penjualan":xlsx("penjualan.xlsx",data["sales_headers"],data["sales"]),
      "keuangan":xlsx("keuangan.xlsx",data["finance_headers"],data["finance"]),
      "riwayat_pembelian":xlsx("riwayat_pembelian.xlsx",data["purchase_headers"],data["purchases"]),
      "catalog":xlsx("catalog.xlsx",data["catalog_headers"],data["catalog"]),
      "laporan_operasional":docx("laporan_operasional.docx","Laporan Operasional",data.get("report",[]))
    }